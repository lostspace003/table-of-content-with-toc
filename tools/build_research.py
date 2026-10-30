"""Per-course research and content plan: one document per course, built from tools/research/<course>.json.
Outputs in TOC-output/<course>/Research/: Markdown, HTML, Word (.docx) and a Gennoor-watermarked PDF.
Also writes the index TOC-output/Research/README.md, grouped by domain."""
import html
import json
import re
import asyncio
from pathlib import Path

import docx
from docx.shared import Pt, Mm, RGBColor

from content import COURSES, COMMON_ASSESSMENT

HERE = Path(__file__).resolve().parent
OUT = HERE.parent / "TOC-output"
SRC = HERE / "research"
DATE = "30 October 2026"
DOMAINS = ["Cybersecurity", "Network and Infrastructure Automation", "Business Process Automation"]
CITE = re.compile(r"\[(c\d+(?:\s*,\s*c\d+)*)\]")


# --- document model: a flat list of blocks rendered three ways -----------------------------------
# ("h", level, text) ("p", text) ("ul", [text]) ("ol", [text]) ("table", header, rows) ("quote", text) ("kv", label, text)

def plan_blocks(c, d, num):
    def n(ids):
        return ", ".join(str(num[i]) for i in ids if i in num)

    B = [("h", 1, f"{c['title']}: Research and Content Plan"),
         ("p", f"{d['domain']} · {c['days']} days (8 hours per day) · {c['level']} · Compiled {DATE}"),
         ("p", "This document is the working plan for building the course from its outline. It records where the course came from, "
               "why the outline is shaped the way it is, what goes into every module (lessons, teaching points, demo, lab, pitfalls and "
               "knowledge checks), the lab environment to prepare, and the sources behind each fact. Numbers in square brackets point to "
               "the reference list at the end; every URL there was opened and checked against the fact it supports.")]

    B += [("h", 2, "1. Where this course came from"), ("kv", "Client input", ""), ("quote", d["source_input"])]
    B += [("p", t) for t in d["why_this_course"]]

    B += [("h", 2, "2. How the outline was shaped"),
          ("table", ["Decision", "Rationale"], [[x["decision"], x["rationale"]] for x in d["design_decisions"]])]

    B += [("h", 2, "3. Audience and prerequisites"),
          ("table", ["Role", "What they need", "Starting point"],
           [[x["role"], x["needs"], x["starting_point"]] for x in d["audience_profile"]]),
          ("h", 3, "Prerequisites from the outline"), ("ul", c["prerequisites"]),
          ("h", 3, "Checks before enrolment"), ("ul", d["prerequisite_check"]),
          ("h", 3, "Course outcomes"), ("ul", c["outcomes"])]

    lab = d["lab_environment"]
    B += [("h", 2, "4. Lab environment"), ("p", lab["overview"]),
          ("h", 3, "Components"),
          ("table", ["Component", "Version or edition", "Purpose", "How to obtain", "Notes"],
           [[x["name"], x["version"], x["purpose"], x["how_to_obtain"], x.get("notes", "")] for x in lab["components"]]),
          ("h", 3, "Fixtures and sample data"),
          ("table", ["Fixture", "Contents", "Used in"], [[x["name"], x["contents"], x["used_in"]] for x in lab["fixtures"]]),
          ("h", 3, "Lab safety rules"), ("ul", lab["safety"])]

    B += [("h", 2, "5. Module-by-module content plan")]
    day = None
    for i, (m, om) in enumerate(zip(d["modules"], c["modules"]), 1):
        if m["day"] != day:
            day = m["day"]; B.append(("h", 3, f"Day {day}"))
        B += [("h", 4, f"Module {i}: {m['title']}"), ("kv", "Goal", m["goal"]),
              ("kv", "Outline topics", ""), ("ul", om["topics"]),
              ("kv", "Learning objectives", ""), ("ul", m["learning_objectives"])]
        for ls in m["lessons"]:
            B += [("h", 5, ls["title"]), ("ul", ls["points"])]
        B += [("kv", "Instructor demo", ""), ("ol", m["instructor_demo"]),
              ("kv", "Hands-on lab", m["lab"]["brief"]), ("ol", m["lab"]["steps"]),
              ("kv", "Lab success criteria", ""), ("ul", m["lab"]["success_criteria"]),
              ("kv", "Deliverable", om["deliverable"]),
              ("kv", "Common pitfalls", ""), ("ul", m["common_pitfalls"]),
              ("kv", "Knowledge check", ""), ("ol", [f"{q['q']} Answer: {q['a']}" for q in m["knowledge_check"]]),
              ("kv", "Slides and assets to prepare", ""), ("ul", m["slides_and_assets"])]
        if m.get("sources"):
            B.append(("kv", "Sources for this module", n(m["sources"])))

    cp = d["capstone_plan"]
    B += [("h", 2, "6. Capstone plan"), ("p", c["capstone"]), ("p", cp["brief"]),
          ("h", 3, "Milestones"), ("ol", cp["milestones"]),
          ("h", 3, "Evaluation notes"), ("ul", cp["evaluation_notes"]),
          ("h", 3, "Assessment weights"), ("table", ["Assessment area", "Weight", "Evidence"], [list(r) for r in COMMON_ASSESSMENT])]

    pv = d["promo_video_research"]
    B += [("h", 2, "7. Promotional video: research behind the script"), ("kv", "Opening scenario", pv["hook_basis"]),
          ("table", ["Claim or on-screen text", "Basis"], [[x["claim"], x["basis"]] for x in pv["claims_check"]])]

    B += [("h", 2, "References")]
    B += [("ref", num[x["id"]], x) for x in d["citations"]]
    B += [("p", "For more info - https://gennoor.com")]
    return B


def numbering(d):
    return {x["id"]: i for i, x in enumerate(d["citations"], 1)}


def cites(text, num, fmt):
    def rep(mo):
        ids = [s.strip() for s in mo.group(1).split(",")]
        return fmt([num[i] for i in ids if i in num])
    return CITE.sub(rep, text)


# --- Markdown ------------------------------------------------------------------------------------

def to_md(B, num):
    f = lambda t: cites(t, num, lambda ns: "[" + ", ".join(map(str, ns)) + "]")
    cell = lambda t: f(str(t)).replace("|", "\\|").replace("\n", " ")
    out = []
    for b in B:
        k = b[0]
        if k == "h": out.append("#" * b[1] + " " + b[2] + "\n")
        elif k == "p": out.append(f(b[1]) + "\n")
        elif k == "quote": out.append("> " + f(b[1]) + "\n")
        elif k == "kv": out.append(f"**{b[1]}.**" + (" " + f(b[2]) if b[2] else "") + "\n")
        elif k == "ul": out.append("\n".join("- " + f(x) for x in b[1]) + "\n")
        elif k == "ol": out.append("\n".join(f"{i}. " + f(x) for i, x in enumerate(b[1], 1)) + "\n")
        elif k == "table":
            out.append("| " + " | ".join(b[1]) + " |\n|" + "---|" * len(b[1]))
            out.append("\n".join("| " + " | ".join(cell(v) for v in r) + " |" for r in b[2]) + "\n")
        elif k == "ref":
            x = b[2]; out.append(f"[{b[1]}] {x['title']}. {x['publisher']}. {x['url']}  \n")
    return "\n".join(out)


# --- HTML ----------------------------------------------------------------------------------------

CSS = """:root{--ink:#111827;--ink2:#4b5563;--line:#e5e7eb;--soft:#f3f6fb;--acc:#2563EB;--bg:#fff}
@media (prefers-color-scheme:dark){:root:not([data-theme=light]){--ink:#e5e7eb;--ink2:#9ca3af;--line:#374151;--soft:#1f2937;--acc:#60a5fa;--bg:#0b1220}}
:root[data-theme=dark]{--ink:#e5e7eb;--ink2:#9ca3af;--line:#374151;--soft:#1f2937;--acc:#60a5fa;--bg:#0b1220}
body{margin:0;background:var(--bg);color:var(--ink);font-family:Inter,"Segoe UI",Helvetica,Arial,sans-serif;line-height:1.55;font-size:14px}
main{max-width:920px;margin:0 auto;padding:28px 16px 48px}
h1{font-size:28px;line-height:1.2;margin:0 0 6px;border-bottom:3px solid var(--acc);padding-bottom:10px}
h2{font-size:20px;margin:34px 0 10px;border-bottom:1px solid var(--line);padding-bottom:5px}
h3{font-size:16px;margin:22px 0 8px;color:var(--acc)}h4{font-size:16px;margin:26px 0 6px;padding:8px 12px;background:var(--soft);border-left:4px solid var(--acc);border-radius:4px}
h5{font-size:14px;margin:14px 0 4px}li{margin:4px 0}p{margin:8px 0}.meta{color:var(--ink2)}
blockquote{margin:8px 0;padding:10px 14px;border-left:4px solid var(--line);background:var(--soft);color:var(--ink2)}
.tw{overflow-x:auto}table{border-collapse:collapse;width:100%;margin:8px 0;font-size:13px}th,td{border:1px solid var(--line);padding:6px 8px;text-align:left;vertical-align:top}th{background:var(--soft)}
a{color:var(--acc);text-decoration:none}sup a{font-size:11px}.ref{font-size:12.5px;word-break:break-word;margin:5px 0}.ref b{color:var(--acc)}
.toc{columns:2;font-size:13px}@media(max-width:640px){.toc{columns:1}}
@page{size:A4;margin:14mm 12mm}@media print{h4{break-after:avoid}h2{break-before:auto}table,blockquote{break-inside:avoid}}"""


def to_html(B, num, title):
    def f(t):
        t = html.escape(str(t))
        return cites(t, num, lambda ns: "<sup>[" + ",".join(f"<a href='#r{n}'>{n}</a>" for n in ns) + "]</sup>")
    toc = [b for b in B if b[0] == "h" and b[1] == 2]
    out = [f"<!doctype html><html lang='en'><head><meta charset='utf-8'><meta name='viewport' content='width=device-width,initial-scale=1'>"
           f"<title>{html.escape(title)}</title><link href='https://fonts.googleapis.com/css2?family=Inter:wght@400;600;800&display=swap' rel='stylesheet'>"
           f"<style>{CSS}</style></head><body><main>"]
    for j, b in enumerate(B):
        k = b[0]
        if k == "h":
            aid = f" id='s{toc.index(b)}'" if b in toc else ""
            out.append(f"<h{b[1]}{aid}>{html.escape(b[2])}</h{b[1]}>")
        elif k == "p":
            cls = " class='meta'" if j == 1 else ""
            out.append(f"<p{cls}>{f(b[1])}</p>")
            if j == 2:
                out.append("<ul class='toc'>" + "".join(f"<li><a href='#s{i}'>{html.escape(h[2])}</a></li>" for i, h in enumerate(toc)) + "</ul>")
        elif k == "quote": out.append(f"<blockquote>{f(b[1])}</blockquote>")
        elif k == "kv": out.append(f"<p><b>{html.escape(b[1])}.</b> {f(b[2])}</p>")
        elif k == "ul": out.append("<ul>" + "".join(f"<li>{f(x)}</li>" for x in b[1]) + "</ul>")
        elif k == "ol": out.append("<ol>" + "".join(f"<li>{f(x)}</li>" for x in b[1]) + "</ol>")
        elif k == "table":
            out.append("<div class='tw'><table><tr>" + "".join(f"<th>{html.escape(h)}</th>" for h in b[1]) + "</tr>"
                       + "".join("<tr>" + "".join(f"<td>{f(v)}</td>" for v in r) + "</tr>" for r in b[2]) + "</table></div>")
        elif k == "ref":
            x = b[2]; u = html.escape(x["url"])
            out.append(f"<p class='ref' id='r{b[1]}'><b>[{b[1]}]</b> {html.escape(x['title'])}. {html.escape(x['publisher'])}. "
                       f"<a href='{u}'>{u}</a><br><span class='meta'>Supports: {html.escape(x['supports'])}. {html.escape(x['key_fact'])}</span></p>")
    out.append("</main></body></html>")
    return "\n".join(out)


# --- Word ----------------------------------------------------------------------------------------

def to_docx(B, num, title, path):
    f = lambda t: cites(str(t), num, lambda ns: "[" + ", ".join(map(str, ns)) + "]")
    doc = docx.Document()
    st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10.5); st.font.color.rgb = RGBColor(0, 0, 0)
    for nm in ["Title", "Heading 1", "Heading 2", "Heading 3", "Heading 4"]:
        doc.styles[nm].font.color.rgb = RGBColor(0, 0, 0); doc.styles[nm].font.name = "Calibri"
    sec = doc.sections[0]; sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin = sec.right_margin = Mm(18); sec.top_margin = sec.bottom_margin = Mm(16)
    for b in B:
        k = b[0]
        if k == "h": doc.add_heading(b[2], 0 if b[1] == 1 else min(b[1] - 1, 4))
        elif k == "p": doc.add_paragraph(f(b[1]))
        elif k == "quote":
            p = doc.add_paragraph(style="Intense Quote"); p.add_run(f(b[1]))
        elif k == "kv":
            p = doc.add_paragraph(); p.add_run(b[1] + ". ").bold = True
            if b[2]: p.add_run(f(b[2]))
        elif k in ("ul", "ol"):
            for x in b[1]:
                doc.add_paragraph(f(x), style="List Bullet" if k == "ul" else "List Number")
        elif k == "table":
            t = doc.add_table(rows=1, cols=len(b[1])); t.style = "Table Grid"
            for i, h in enumerate(b[1]):
                c = t.rows[0].cells[i]; c.text = ""; c.paragraphs[0].add_run(h).bold = True
            for r in b[2]:
                cells = t.add_row().cells
                for i, v in enumerate(r):
                    cells[i].text = f(v)
            doc.add_paragraph()
        elif k == "ref":
            x = b[2]; p = doc.add_paragraph(); p.add_run(f"[{b[1]}] ").bold = True
            p.add_run(f"{x['title']}. {x['publisher']}. {x['url']}")
    doc.core_properties.author = "Gennoor Tech"; doc.core_properties.title = title
    doc.save(str(path))


async def html_to_pdf(pairs):
    from playwright.async_api import async_playwright
    async with async_playwright() as p:
        b = await p.chromium.launch(executable_path="/opt/pw-browsers/chromium")
        pg = await b.new_page()
        for h, pdf in pairs:
            await pg.goto(h.as_uri()); await pg.wait_for_timeout(400)
            await pg.pdf(path=str(pdf), format="A4", print_background=True,
                         margin={"top": "14mm", "bottom": "14mm", "left": "12mm", "right": "12mm"})
        await b.close()


def check(c, d):
    assert [m["title"] for m in d["modules"]] == [m["title"] for m in c["modules"]], c["title"] + ": module titles differ"
    ids = {x["id"] for x in d["citations"]}
    missing = set(re.findall(r"c\d+", " ".join(CITE.findall(json.dumps(d))))) - ids
    assert not missing, f"{c['title']}: unknown citations {missing}"
    assert d["domain"] in DOMAINS, d["domain"]


def main():
    import gennoor_watermark as wm
    built, pdfs = [], []
    for c in COURSES:
        src = SRC / f"{c['title']}.json"
        if not src.exists():
            print("skip (no research file):", c["title"]); continue
        d = json.loads(src.read_text()); check(c, d)
        num = numbering(d); B = plan_blocks(c, d, num)
        out = OUT / c["folder"] / "Research"; out.mkdir(parents=True, exist_ok=True)
        base = f"{c['title']} - Research and Content Plan"
        (out / f"{base}.md").write_text(to_md(B, num))
        h = out / f"{base}.html"; h.write_text(to_html(B, num, base))
        to_docx(B, num, base, out / f"{base}.docx")
        pdfs.append((h, out / "~plain.pdf", out / f"{base} - Gennoor.pdf"))
        built.append((c, d, base))
    asyncio.run(html_to_pdf([(h, t) for h, t, _ in pdfs]))
    for _, tmp, final in pdfs:
        final.write_bytes(wm.stamp_bytes(tmp.read_bytes(), mode="overlay")); tmp.unlink()
    write_index(built)
    for c, d, _ in built:
        print(f"built {c['title']}: {len(d['modules'])} modules, {len(d['citations'])} citations")


def write_index(built):
    idx = OUT / "Research"; idx.mkdir(exist_ok=True)
    md = ["# Research and Content Plans\n",
          "Each course has its own research document in `TOC-output/<course>/Research/` (watermarked PDF, Word, Markdown and HTML). "
          "Each one covers where the course came from, how the outline was shaped, the lab environment, a module-by-module content plan "
          "(lessons, teaching points, demo, lab steps, pitfalls, knowledge checks, assets to prepare), the capstone, the research behind "
          "the promo video and a verified reference list.\n"]
    for dom in DOMAINS:
        rows = [(c, d, b) for c, d, b in built if d["domain"] == dom]
        if not rows: continue
        md.append(f"\n## {dom}\n\n| Course | Days | Modules | Sources | Document |\n|---|---|---|---|---|")
        for c, d, b in rows:
            p = f"../{c['folder']}/Research/{b} - Gennoor.pdf".replace(" ", "%20")
            md.append(f"| {c['title']} | {c['days']} | {len(d['modules'])} | {len(d['citations'])} | [PDF]({p}) |")
    md.append("\nRebuild with `python tools/build_research.py` (data in `tools/research/<course>.json`, format in `tools/research/SCHEMA.md`).\n")
    (idx / "README.md").write_text("\n".join(md))


if __name__ == "__main__":
    main()
