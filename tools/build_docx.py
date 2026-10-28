"""Plain black-and-white Word (.docx) version of every course outline."""
from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Mm, RGBColor

from content import COURSES, COMMON_ASSESSMENT, SHARED_PRINCIPLES

HERE = Path(__file__).resolve().parent
OUT = HERE.parent / "TOC-output"
BLACK = RGBColor(0, 0, 0)


def _plain(doc):
    st = doc.styles["Normal"]
    st.font.name = "Calibri"; st.font.size = Pt(11); st.font.color.rgb = BLACK
    for name in ["Heading 1", "Heading 2", "Heading 3", "Title"]:
        s = doc.styles[name]; s.font.color.rgb = BLACK; s.font.name = "Calibri"
    doc.styles["Title"].font.size = Pt(22)
    doc.styles["Heading 1"].font.size = Pt(16)
    doc.styles["Heading 2"].font.size = Pt(13)
    doc.styles["Heading 3"].font.size = Pt(11.5)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin = sec.right_margin = Mm(20); sec.top_margin = sec.bottom_margin = Mm(18)


def _table(doc, header, rows):
    t = doc.add_table(rows=1, cols=len(header)); t.style = "Table Grid"
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]; c.text = ""; r = c.paragraphs[0].add_run(h); r.bold = True
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = str(v)
    doc.add_paragraph()


def build(c):
    doc = docx.Document(); _plain(doc)
    doc.add_paragraph("Gennoor Tech - Course Outline").alignment = WD_ALIGN_PARAGRAPH.RIGHT
    doc.add_heading(c["title"], 0)
    doc.add_paragraph(c["tagline"])
    _table(doc, ["Duration", "Level", "Format"],
           [[f"{c['days']} days (8 hours per day)", c["level"], "Instructor-led, hands-on labs, capstone project"]])
    doc.add_heading("Who should attend", 1); doc.add_paragraph(c["audience"] + ".")
    doc.add_heading("Prerequisites", 1)
    for p in c["prerequisites"]:
        doc.add_paragraph(p, style="List Bullet")
    doc.add_heading("What you will be able to do", 1)
    for o in c["outcomes"]:
        doc.add_paragraph(o, style="List Bullet")
    doc.add_heading("Tools and technologies", 1)
    doc.add_paragraph(", ".join(label for _, label in c["tools"]) + ".")
    doc.add_heading("Table of contents", 1)
    days = {}
    for m in c["modules"]:
        days.setdefault(m["day"], []).append(m)
    n = 0
    for d in sorted(days):
        doc.add_heading(f"Day {d}", 2)
        for m in days[d]:
            n += 1
            doc.add_heading(f"Module {n}: {m['title']}", 3)
            for t in m["topics"]:
                doc.add_paragraph(t, style="List Bullet")
            p = doc.add_paragraph(); p.add_run("Hands-on lab: ").bold = True; p.add_run(m["lab"])
            p = doc.add_paragraph(); p.add_run("Deliverable: ").bold = True; p.add_run(m["deliverable"])
    doc.add_heading("Capstone project", 1); doc.add_paragraph(c["capstone"])
    doc.add_heading("Assessment", 1)
    _table(doc, ["Assessment area", "Weight", "Evidence"], COMMON_ASSESSMENT)
    doc.add_heading("Engineering principles used throughout", 1)
    for p in SHARED_PRINCIPLES:
        doc.add_paragraph(p, style="List Bullet")
    doc.add_paragraph("For more info - https://gennoor.com")
    d = OUT / c["folder"] / "TOC"; d.mkdir(parents=True, exist_ok=True)
    out = d / f"{c['title']} - Course Outline.docx"
    doc.core_properties.author = "Gennoor Tech"; doc.core_properties.title = f"{c['title']} - Course Outline"
    doc.save(str(out)); return out


if __name__ == "__main__":
    for c in COURSES:
        print("built", build(c).name)
